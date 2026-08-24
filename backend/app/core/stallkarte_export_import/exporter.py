import datetime
from copy import copy
from decimal import Decimal
from typing import Iterable, cast

import openpyxl
from openpyxl import Workbook
from openpyxl.styles.fonts import Font
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string, quote_sheetname
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from app import domain
from app.core.stallkarteexport.mapping import (
    OutdoorJournalDaysMapping,
    ThroughputReportGeneralNotesMapping,
    QualityReportBioChicksMapping,
    YesNoMapping,
    default_stallkarte_checklist_mapping,
    default_export_worksheet_names,
    default_stallkarte_holding_mapping,
    default_outdoor_journal_days_mapping,
    default_outdoor_journal_header_mapping,
    default_quality_report_bio_chicks_mapping,
    default_throughput_report_header_mapping,
    default_throughput_report_general_notes_mapping,
    default_throughput_report_finish_notes_mapping,
    default_throughput_report_mortality_mapping,
    get_mapping_of_production_day,
    get_throughput_report_mortality_row_pairs,
)
from app.domain.stallkarte.events import models as event_models
from app.domain import StallkarteStateDay
from app.domain.stallkarte.events.models import NoteEntry, NoteType


def format_date(date: datetime.date) -> str:
    return date.strftime("%d.%m.%Y")


def optional_sum(a: int | None, b: int | None) -> int | None:
    if a is None and b is None:
        return None
    return (a if a is not None else 0) + (b if b is not None else 0)


def fill_yes_no_cell(ws: Worksheet, mapping: YesNoMapping, value: bool | None) -> None:
    yes_ref = resolve_cell_reference(ws, mapping.yes)
    no_ref = resolve_cell_reference(ws, mapping.no)

    if value:
        set_cell_value(ws, mapping.yes, "Ja")
        ws[yes_ref].font = Font(bold=True)
        set_cell_value(ws, mapping.no, "")
    else:
        set_cell_value(ws, mapping.yes, "")
        set_cell_value(ws, mapping.no, "Nein")
        ws[no_ref].font = Font(bold=True)


def build_sheet_sum_formula(sheet_names: list[str], cell: str) -> str:
    if not sheet_names:
        return "=0"

    refs = ",".join(
        f"{quote_sheetname(sheet_name)}!{cell}" for sheet_name in sheet_names
    )
    return f"=SUM({refs})"


type CellValue = (
    str | int | float | bool | datetime.date | datetime.datetime | Decimal | None
)


def resolve_cell_reference(ws: Worksheet, cell: str) -> str:
    try:
        merged_ranges = cast(Iterable[CellRange], ws.merged_cells.ranges)
    except Exception:
        return cell

    for merged_range in merged_ranges:
        if cell in merged_range:
            return f"{get_column_letter(merged_range.min_col)}{merged_range.min_row}"

    return cell


def set_cell_value(ws: Worksheet, cell: str, value: CellValue) -> None:
    ws[resolve_cell_reference(ws, cell)] = value


def copy_cell_style(ws: Worksheet, source_cell: str, target_cell: str) -> None:
    source_ref = resolve_cell_reference(ws, source_cell)
    target_ref = resolve_cell_reference(ws, target_cell)

    ws[target_ref]._style = copy(ws[source_ref]._style)


def ensure_merged_like_source(
    ws: Worksheet, source_cell: str, target_cell: str
) -> None:
    try:
        merged_ranges = cast(Iterable[CellRange], ws.merged_cells.ranges)
    except Exception:
        return

    source_ref = resolve_cell_reference(ws, source_cell)
    source_range = next((rng for rng in merged_ranges if source_ref in rng), None)
    if source_range is None:
        return

    target_col, target_row = coordinate_from_string(target_cell)
    target_min_col = column_index_from_string(target_col)

    source_start, source_end = str(source_range.coord).split(":")
    source_start_col, source_min_row = coordinate_from_string(source_start)
    source_end_col, source_max_row = coordinate_from_string(source_end)

    source_min_col = column_index_from_string(source_start_col)
    source_max_col = column_index_from_string(source_end_col)

    col_span = source_max_col - source_min_col
    row_span = source_max_row - source_min_row

    target_min_row = target_row
    target_max_col = target_min_col + col_span
    target_max_row = target_min_row + row_span

    target_range = CellRange(
        min_col=target_min_col,
        min_row=target_min_row,
        max_col=target_max_col,
        max_row=target_max_row,
    )

    if any(rng.coord == target_range.coord for rng in merged_ranges):
        return

    source_style = copy(ws[source_ref]._style)
    for row in range(target_min_row, target_max_row + 1):
        for col in range(target_min_col, target_max_col + 1):
            cell_ref = f"{get_column_letter(col)}{row}"
            ws[cell_ref]._style = copy(source_style)

    try:
        ws.merge_cells(target_range.coord)
    except ValueError:
        return


def section_number_by_index(state: domain.StallkarteState, index: int) -> int | None:
    rearing_section_number = (
        state.rearing_farm.sections[index].number
        if state.rearing_farm is not None and index < len(state.rearing_farm.sections)
        else None
    )
    fattening_section_number = (
        state.fattening_farm.sections[index].number
        if state.fattening_farm is not None
        and index < len(state.fattening_farm.sections)
        else None
    )
    return rearing_section_number or fattening_section_number


def section_name_by_index(state: domain.StallkarteState, index: int) -> str | None:
    rearing_section_name = (
        state.rearing_farm.sections[index].name
        if state.rearing_farm is not None and index < len(state.rearing_farm.sections)
        else None
    )
    fattening_section_name = (
        state.fattening_farm.sections[index].name
        if state.fattening_farm is not None
        and index < len(state.fattening_farm.sections)
        else None
    )
    name = rearing_section_name or fattening_section_name
    return name if isinstance(name, str) and name != "" else None


def get_section_from_farm(
    farm: domain.StallkarteFarm,
    section_number: int,
) -> domain.StallkarteSection | None:
    for section in farm.sections:
        if section.number == section_number:
            return section

    return None


def get_stall_section_numbers(state: domain.StallkarteState) -> list[int]:
    if state.rearing_farm is not None:
        return [section.number for section in state.rearing_farm.sections]

    if state.fattening_farm is not None:
        return [section.number for section in state.fattening_farm.sections]

    return []


def get_installation_details_by_section(
    state: domain.StallkarteState,
) -> dict[int, event_models.SectionInstallationDetails]:
    details = state.installation_details_by_section
    if isinstance(details, dict):
        return details
    return {}


def format_parent_flocks_for_section(
    state: domain.StallkarteState,
    section_number: int,
) -> str:
    details = get_installation_details_by_section(state).get(section_number)
    if details is None:
        return ""

    identifiers = [
        parent_flock.herd_identifier
        for parent_flock in details.parent_flocks
        if parent_flock.herd_identifier != ""
    ]
    return "; ".join(identifiers)


def format_production_weeks_for_section(
    state: domain.StallkarteState,
    section_number: int,
) -> str:
    details = get_installation_details_by_section(state).get(section_number)
    if details is None:
        return ""

    production_weeks = [
        str(parent_flock.production_week) for parent_flock in details.parent_flocks
    ]
    return "; ".join(production_weeks)


def get_initial_animals_count_for_section(
    state: domain.StallkarteState,
    section_number: int,
) -> int | None:
    details = get_installation_details_by_section(state).get(section_number)
    if details is None:
        if state.transfer is not None:
            return state.transfer.animals_by_section_number.get(section_number)
        return None
    return details.initial_animals_count


def get_bedding(state: domain.StallkarteState) -> str:
    section_numbers = get_stall_section_numbers(state)
    for section_number in section_numbers:
        details = get_installation_details_by_section(state).get(section_number)
        if details is not None and details.bedding.strip() != "":
            return details.bedding
    return ""


def format_parent_flocks_by_section(state: domain.StallkarteState) -> str:
    section_values: list[str] = []
    for section_number in get_stall_section_numbers(state):
        parent_flock = format_parent_flocks_for_section(state, section_number)
        if parent_flock == "":
            continue
        section_values.append(f"{section_number}: {parent_flock}")

    if section_values:
        return " | ".join(section_values)

    return ""


def format_production_weeks_by_section(state: domain.StallkarteState) -> str:
    section_values: list[str] = []
    for section_number in get_stall_section_numbers(state):
        production_week = format_production_weeks_for_section(state, section_number)
        if production_week == "":
            continue
        section_values.append(f"{section_number}: {production_week}")

    if section_values:
        return " | ".join(section_values)

    return ""


def format_note_event_type(note_type: event_models.NoteType) -> str:
    return note_type.to_label()


def format_value_and_unit(
    value: float | int | None,
    unit: str | None,
) -> str | None:
    if value is None and unit is None:
        return None

    parts: list[str] = []
    if value is not None:
        parts.append(str(value))
    if unit is not None:
        parts.append(unit)

    return " ".join(parts)


def format_value_and_unit_if_complete(
    value: float | int | None,
    unit: str | None,
) -> str | None:
    if value is None or unit is None:
        return None

    return f"{value} {unit}"


def waiting_time_to_days(note: event_models.NoteEntry) -> int | None:
    if (
        note.treatment_waiting_time_value is None
        or note.treatment_waiting_time_unit is None
    ):
        return None

    if note.treatment_waiting_time_unit == event_models.WaitingTimeUnit.DAY:
        return note.treatment_waiting_time_value

    if note.treatment_waiting_time_unit == event_models.WaitingTimeUnit.WEEK:
        return note.treatment_waiting_time_value * 7

    return None


def iter_general_notes_by_type(
    stallkarte: domain.Stallkarte,
    note_type: NoteType,
) -> list[tuple[int, event_models.NoteEntry]]:
    notes: list[tuple[int, event_models.NoteEntry]] = []

    for production_day in sorted(stallkarte.state.days.keys()):
        day = stallkarte.state.days[production_day]
        for note in day.notes:
            if note.note_type == note_type:
                notes.append((production_day, note))

    return notes


def format_treatment_note_parts(note: event_models.NoteEntry) -> list[str]:
    parts: list[str] = []

    if note.treatment_code is not None:
        parts.append(note.treatment_code.value)

    amount = format_value_and_unit(
        note.treatment_amount_value,
        note.treatment_amount_unit.value if note.treatment_amount_unit else None,
    )
    if amount is not None:
        parts.append(amount)

    waiting = format_value_and_unit(
        note.treatment_waiting_time_value,
        note.treatment_waiting_time_unit.value
        if note.treatment_waiting_time_unit
        else None,
    )
    if waiting is not None:
        parts.append(f"Wartezeit {waiting}")

    return parts


def format_note_specific_parts(note: event_models.NoteEntry) -> list[str]:
    if note.note_type == event_models.NoteType.VACCINATION:
        return [note.vaccination_code.value.upper()] if note.vaccination_code else []

    if note.note_type == event_models.NoteType.TREATMENT:
        return format_treatment_note_parts(note)

    if note.note_type == event_models.NoteType.SOCK_TEST:
        return [note.sock_test_result.value] if note.sock_test_result else []

    return []


def format_general_note_entry(note: event_models.NoteEntry) -> str:
    parts = [format_note_event_type(note.note_type)]
    parts.extend(format_note_specific_parts(note))

    note_text = note.note_text.strip() if note.note_text is not None else ""
    if note_text:
        parts.append(note_text)

    return " - ".join(parts)


def get_quality_report_stall_columns(
    mp: QualityReportBioChicksMapping,
    section_count: int,
) -> list[tuple[str, str, str]]:
    total_stalls = max(len(mp.stall_columns), section_count)
    first_stall_first_col = column_index_from_string(mp.stall_columns[0][0])
    stall_width = len(mp.stall_columns[0])

    columns: list[tuple[str, str, str]] = []
    for stall_index in range(total_stalls):
        base = first_stall_first_col + (stall_index * stall_width)
        columns.append(
            (
                get_column_letter(base),
                get_column_letter(base + 1),
                get_column_letter(base + 2),
            )
        )

    return columns


def fill_quality_report_stall_headings(
    ws: Worksheet,
    mp: QualityReportBioChicksMapping,
    state: domain.StallkarteState,
    section_count: int,
) -> None:
    stalls = get_quality_report_stall_columns(mp, section_count)
    for stall_index, columns in enumerate(stalls):
        if stall_index >= section_count:
            set_cell_value(ws, f"{columns[0]}{mp.stall_heading_row}", "")
            continue

        section_number = section_number_by_index(state, stall_index)
        heading_cell = f"{columns[0]}{mp.stall_heading_row}"
        heading_source_cell = f"{stalls[0][0]}{mp.stall_heading_row}"

        ensure_merged_like_source(ws, heading_source_cell, heading_cell)
        copy_cell_style(ws, heading_source_cell, heading_cell)
        set_cell_value(
            ws,
            heading_cell,
            (
                f"{mp.stall_heading_prefix} {section_number}"
                if section_number is not None
                else ""
            ),
        )


def fill_quality_report_stall_subheadings(
    ws: Worksheet,
    mp: QualityReportBioChicksMapping,
    section_count: int,
) -> None:
    stalls = get_quality_report_stall_columns(mp, section_count)
    source_columns = stalls[0]

    for stall_index, columns in enumerate(stalls):
        if stall_index >= section_count:
            for column in columns:
                set_cell_value(ws, f"{column}{mp.stall_subheading_row}", "")
            continue

        for index, column in enumerate(columns):
            source_cell = f"{source_columns[index]}{mp.stall_subheading_row}"
            target_cell = f"{column}{mp.stall_subheading_row}"

            ensure_merged_like_source(ws, source_cell, target_cell)
            copy_cell_style(ws, source_cell, target_cell)
            source_ref = resolve_cell_reference(ws, source_cell)
            set_cell_value(ws, target_cell, ws[source_ref].value)


def hide_unused_quality_report_stalls(
    ws: Worksheet,
    mp: QualityReportBioChicksMapping,
    section_count: int,
) -> None:
    if not hasattr(ws, "column_dimensions"):
        return

    stalls = get_quality_report_stall_columns(mp, section_count)
    for stall_index, columns in enumerate(stalls):
        hidden = stall_index >= section_count
        for column in columns:
            ws.column_dimensions[column].hidden = hidden


def fill_quality_report_mortality_grid(
    ws: Worksheet,
    mp: QualityReportBioChicksMapping,
    section_sheet_names: list[str],
) -> None:
    stalls = get_quality_report_stall_columns(mp, len(section_sheet_names))

    for stall_index, columns in enumerate(stalls):
        sheet_name = (
            section_sheet_names[stall_index]
            if stall_index < len(section_sheet_names)
            else None
        )

        for report_row, template_row in zip(
            range(mp.mortality_first_row, mp.mortality_last_row + 1),
            range(
                mp.template_first_row,
                mp.template_first_row
                + (mp.mortality_last_row - mp.mortality_first_row + 1),
            ),
            strict=True,
        ):
            source_columns = stalls[0]
            for index, column in enumerate(columns):
                copy_cell_style(
                    ws,
                    f"{source_columns[index]}{report_row}",
                    f"{column}{report_row}",
                )

            if sheet_name is None:
                for column in columns:
                    set_cell_value(ws, f"{column}{report_row}", "")
                continue

            sheet_ref = quote_sheetname(sheet_name)
            set_cell_value(
                ws, f"{columns[0]}{report_row}", f"={sheet_ref}!C{template_row}"
            )
            set_cell_value(
                ws, f"{columns[1]}{report_row}", f"={sheet_ref}!D{template_row}"
            )
            set_cell_value(
                ws, f"{columns[2]}{report_row}", f"={sheet_ref}!E{template_row}"
            )

        summary_row = mp.summary_row
        source_columns = stalls[0]
        for index, column in enumerate(columns):
            copy_cell_style(
                ws,
                f"{source_columns[index]}{summary_row}",
                f"{column}{summary_row}",
            )

        if sheet_name is None:
            for column in columns:
                set_cell_value(ws, f"{column}{summary_row}", "")
            continue

        set_cell_value(
            ws,
            f"{columns[0]}{summary_row}",
            f"=SUM({columns[0]}{mp.mortality_first_row}:{columns[0]}{mp.mortality_last_row})",
        )
        set_cell_value(
            ws,
            f"{columns[1]}{summary_row}",
            f"=SUM({columns[1]}{mp.mortality_first_row}:{columns[1]}{mp.mortality_last_row})",
        )
        set_cell_value(
            ws,
            f"{columns[2]}{summary_row}",
            f"=SUM({columns[2]}{mp.mortality_first_row}:{columns[2]}{mp.mortality_last_row})",
        )


def build_avg_weight_label(template_label: object, section_label: str) -> str:
    if isinstance(template_label, str) and template_label.endswith("- "):
        return f"{template_label}{section_label}"

    if isinstance(template_label, str) and " - " in template_label:
        prefix = template_label.rsplit(" - ", maxsplit=1)[0]
        return f"{prefix} - {section_label}"

    return f"Durchschnittsgewicht am 7. Tag in Gramm - {section_label}"


def copy_quality_report_row_style(
    ws: Worksheet,
    source_row: int,
    target_row: int,
    label_column: str,
) -> None:
    max_column = ws.max_column if isinstance(ws.max_column, int) else 11
    for col_idx in range(1, max_column + 1):
        source_cell = f"{get_column_letter(col_idx)}{source_row}"
        target_cell = f"{get_column_letter(col_idx)}{target_row}"
        copy_cell_style(ws, source_cell, target_cell)

    ensure_merged_like_source(
        ws,
        f"{label_column}{source_row}",
        f"{label_column}{target_row}",
    )

    try:
        if ws.row_dimensions[source_row].height is not None:
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    except AttributeError:
        pass


def shift_merged_ranges_down(
    ws: Worksheet,
    start_row: int,
    offset: int,
) -> None:
    if offset <= 0:
        return

    try:
        merged_ranges = list(cast(Iterable[CellRange], ws.merged_cells.ranges))
    except AttributeError:
        return
    to_shift: list[tuple[str, str, str, CellValue, object]] = []

    for merged_range in merged_ranges:
        min_row = cast(int, merged_range.min_row)
        max_row = cast(int, merged_range.max_row)
        min_col = cast(int, merged_range.min_col)
        max_col = cast(int, merged_range.max_col)

        if min_row >= start_row:
            shifted_min_row = min_row + offset
            top_left_cell = f"{get_column_letter(min_col)}{shifted_min_row}"
            top_left_value = cast(CellValue, ws[top_left_cell].value)
            top_left_style = copy(ws[top_left_cell]._style)

            shifted = CellRange(
                min_col=min_col,
                min_row=shifted_min_row,
                max_col=max_col,
                max_row=max_row + offset,
            )
            to_shift.append(
                (
                    merged_range.coord,
                    shifted.coord,
                    top_left_cell,
                    top_left_value,
                    top_left_style,
                )
            )

    for old_coord, _, _, _, _ in to_shift:
        try:
            ws.unmerge_cells(old_coord)
        except Exception:
            pass

    for _, new_coord, top_left_cell, top_left_value, top_left_style in to_shift:
        ws.merge_cells(new_coord)
        ws[top_left_cell]._style = copy(top_left_style)
        ws[top_left_cell] = top_left_value


def fill_quality_report_avg_weight_day_7(
    ws: Worksheet,
    state: domain.StallkarteState,
    mp: QualityReportBioChicksMapping,
    section_sheet_names: list[str],
) -> None:
    section_count = len(section_sheet_names)
    if section_count == 0:
        return

    label_col = mp.avg_weight_day_7_label_column
    row_template = mp.avg_weight_day_7_template_row
    template_label_cell = f"{label_col}{row_template}"
    template_ref = resolve_cell_reference(ws, template_label_cell)
    template_label = ws[template_ref].value

    if section_count > 1:
        inserted_rows = section_count - 1
        insertion_start = row_template + 1
        max_column = ws.max_column if isinstance(ws.max_column, int) else 11
        max_row = ws.max_row if isinstance(ws.max_row, int) else row_template + 30
        ws.move_range(
            f"A{insertion_start}:{get_column_letter(max_column)}{max_row}",
            rows=inserted_rows,
            cols=0,
        )
        shift_merged_ranges_down(ws, insertion_start, inserted_rows)

    for index in range(section_count):
        row = row_template + index

        if index > 0:
            copy_quality_report_row_style(ws, row_template, row, label_col)

        section_name = section_name_by_index(state, index)
        section_label = section_name or f"Abteil {index + 1}"
        label_value = build_avg_weight_label(template_label, section_label)
        set_cell_value(ws, f"{label_col}{row}", label_value)

    day_7_weight_cell = get_mapping_of_production_day(7).weight

    for index, sheet_name in enumerate(section_sheet_names):
        row = row_template + index
        sheet_ref = quote_sheetname(sheet_name)
        set_cell_value(
            ws,
            f"{mp.avg_weight_day_7_column}{row}",
            f"={sheet_ref}!{day_7_weight_cell}",
        )


def fill_throughput_report_mortality(
    ws: Worksheet,
    section_sheet_names: list[str],
) -> None:
    rearing_pairs, fattening_pairs = get_throughput_report_mortality_row_pairs()
    mp = default_throughput_report_mortality_mapping

    for summary_row, template_row in rearing_pairs:
        set_cell_value(
            ws,
            f"{mp.rearing_column}{summary_row}",
            build_sheet_sum_formula(
                section_sheet_names,
                f"E{template_row}",
            ),
        )

    for summary_row, template_row in fattening_pairs:
        set_cell_value(
            ws,
            f"{mp.fattening_column}{summary_row}",
            build_sheet_sum_formula(
                section_sheet_names,
                f"E{template_row}",
            ),
        )


def fill_throughput_report_total_animals(
    ws: Worksheet,
    section_sheet_names: list[str],
) -> None:
    mp = default_throughput_report_mortality_mapping
    set_cell_value(
        ws,
        mp.total_animals_cell,
        build_sheet_sum_formula(
            section_sheet_names,
            default_stallkarte_holding_mapping.animals_count,
        ),
    )


def fill_throughput_report_header(
    ws: Worksheet,
    holding: domain.AgriculturalHolding,
    stallkarte: domain.Stallkarte,
) -> None:
    state = stallkarte.state
    mp = default_throughput_report_header_mapping

    fattening_stalls = (
        ", ".join(str(section.number) for section in state.fattening_farm.sections)
        if state.fattening_farm is not None
        else ""
    )
    rearing_stalls = (
        ", ".join(str(section.number) for section in state.rearing_farm.sections)
        if state.rearing_farm is not None
        else ""
    )

    set_cell_value(ws, mp.farm_number_cell, "")
    set_cell_value(ws, mp.holding_name_cell, holding.name)
    set_cell_value(
        ws,
        mp.holding_address_cell,
        f"{holding.address_street} {holding.address_zip} {holding.address_city}",
    )
    set_cell_value(ws, mp.fattening_cycle_cell, state.fattening_cycle)

    set_cell_value(
        ws,
        mp.fattening_farm_name_cell,
        state.fattening_farm.name if state.fattening_farm is not None else "",
    )
    set_cell_value(ws, mp.fattening_stalls_cell, fattening_stalls)
    set_cell_value(
        ws,
        mp.fattening_vvvo_number_cell,
        state.fattening_farm.vvvo_number if state.fattening_farm is not None else "",
    )

    set_cell_value(
        ws,
        mp.rearing_farm_name_cell,
        state.rearing_farm.name if state.rearing_farm is not None else "",
    )
    set_cell_value(ws, mp.rearing_stalls_cell, rearing_stalls)
    set_cell_value(
        ws,
        mp.rearing_vvvo_number_cell,
        state.rearing_farm.vvvo_number if state.rearing_farm is not None else "",
    )

    set_cell_value(ws, mp.date_hatched_cell, state.date_hatched)
    set_cell_value(
        ws,
        mp.transfer_date_cell,
        state.transfer.date if state.transfer is not None else "",
    )

    set_cell_value(ws, mp.hatchery_cell, state.hatchery)
    set_cell_value(ws, mp.breed_cell, state.breed)
    section_numbers = get_stall_section_numbers(state)

    set_cell_value(
        ws,
        mp.production_week_cell_primary,
        format_production_weeks_for_section(state, section_numbers[0])
        if len(section_numbers) > 0
        else "",
    )
    set_cell_value(
        ws,
        mp.production_week_cell_secondary,
        format_production_weeks_for_section(state, section_numbers[1])
        if len(section_numbers) > 1
        else "",
    )
    set_cell_value(
        ws,
        mp.production_week_cell_tertiary,
        format_production_weeks_for_section(state, section_numbers[2])
        if len(section_numbers) > 2
        else "",
    )
    set_cell_value(ws, mp.eco_control_number_cell, state.eco_control_number)
    set_cell_value(
        ws,
        mp.parent_flock_cell_primary,
        format_parent_flocks_for_section(state, section_numbers[0])
        if len(section_numbers) > 0
        else "",
    )
    set_cell_value(
        ws,
        mp.parent_flock_cell_secondary,
        format_parent_flocks_for_section(state, section_numbers[1])
        if len(section_numbers) > 1
        else "",
    )
    set_cell_value(
        ws,
        mp.parent_flock_cell_tertiary,
        format_parent_flocks_for_section(state, section_numbers[2])
        if len(section_numbers) > 2
        else "",
    )
    set_cell_value(ws, mp.bedding_cell, get_bedding(state))
    set_cell_value(ws, mp.eu_bio_cell, "X" if state.is_eu_bio else "")
    set_cell_value(ws, mp.naturland_cell, "X" if state.is_naturland else "")


def fill_quality_report_bio_chicks(
    ws: Worksheet,
    holding: domain.AgriculturalHolding,
    stallkarte: domain.Stallkarte,
    section_sheet_names: list[str],
) -> None:
    state = stallkarte.state
    mp: QualityReportBioChicksMapping = default_quality_report_bio_chicks_mapping

    set_cell_value(ws, mp.holding_name_cell, holding.name)
    set_cell_value(
        ws,
        mp.holding_address_cell,
        f"{holding.address_street} {holding.address_zip} {holding.address_city}",
    )
    set_cell_value(
        ws,
        mp.farm_name_cell,
        state.rearing_farm.name
        if state.rearing_farm is not None
        else (state.fattening_farm.name if state.fattening_farm is not None else ""),
    )
    set_cell_value(ws, mp.date_hatched_cell, state.date_hatched)
    set_cell_value(ws, mp.hatchery_cell, state.hatchery)
    set_cell_value(
        ws,
        mp.animals_count_cell,
        build_sheet_sum_formula(
            section_sheet_names,
            default_stallkarte_holding_mapping.animals_count,
        ),
    )

    section_count = len(section_sheet_names)
    fill_quality_report_stall_headings(ws, mp, state, section_count)
    fill_quality_report_stall_subheadings(ws, mp, section_count)
    hide_unused_quality_report_stalls(ws, mp, section_count)

    for index, _ in enumerate(section_sheet_names):
        col_offset = index * mp.farm_section_cellwidth
        section_number = section_number_by_index(state, index)

        # farm section name cell
        farm_section_col = get_column_letter(
            column_index_from_string(mp.farm_section_name_column_start) + col_offset
        )
        ensure_merged_like_source(
            ws,
            f"{mp.farm_section_name_column_start}{mp.farm_section_name_row}",
            f"{farm_section_col}{mp.farm_section_name_row}",
        )
        copy_cell_style(
            ws,
            f"{mp.farm_section_name_column_start}{mp.farm_section_name_row}",
            f"{farm_section_col}{mp.farm_section_name_row}",
        )
        set_cell_value(
            ws,
            f"{farm_section_col}{mp.farm_section_name_row}",
            str(section_number) if section_number is not None else "",
        )

        # parent flock cell
        parent_flock_col = get_column_letter(
            column_index_from_string(mp.parent_flock_column_start) + col_offset
        )
        ensure_merged_like_source(
            ws,
            f"{mp.parent_flock_column_start}{mp.parent_flock_row}",
            f"{parent_flock_col}{mp.parent_flock_row}",
        )
        copy_cell_style(
            ws,
            f"{mp.parent_flock_column_start}{mp.parent_flock_row}",
            f"{parent_flock_col}{mp.parent_flock_row}",
        )
        set_cell_value(
            ws,
            f"{parent_flock_col}{mp.parent_flock_row}",
            format_parent_flocks_for_section(state, section_number)
            if section_number is not None
            else "",
        )

        # production week cell
        production_week_col = get_column_letter(
            column_index_from_string(mp.production_week_column_start) + col_offset
        )
        ensure_merged_like_source(
            ws,
            f"{mp.production_week_column_start}{mp.production_week_row}",
            f"{production_week_col}{mp.production_week_row}",
        )
        copy_cell_style(
            ws,
            f"{mp.production_week_column_start}{mp.production_week_row}",
            f"{production_week_col}{mp.production_week_row}",
        )
        set_cell_value(
            ws,
            f"{production_week_col}{mp.production_week_row}",
            format_production_weeks_for_section(state, section_number)
            if section_number is not None
            else "",
        )

    fill_quality_report_mortality_grid(ws, mp, section_sheet_names)
    fill_quality_report_avg_weight_day_7(ws, state, mp, section_sheet_names)


def fill_outdoor_journal_header(
    ws: Worksheet,
    holding: domain.AgriculturalHolding,
    stallkarte: domain.Stallkarte,
) -> None:
    state = stallkarte.state
    mp = default_outdoor_journal_header_mapping

    stalls = (
        ", ".join(str(section.number) for section in state.rearing_farm.sections)
        if state.rearing_farm is not None
        else (
            ", ".join(str(section.number) for section in state.fattening_farm.sections)
            if state.fattening_farm is not None
            else ""
        )
    )

    set_cell_value(ws, mp.fattening_cycle_cell, state.fattening_cycle)
    set_cell_value(ws, mp.holding_name_cell, holding.name)
    set_cell_value(
        ws,
        mp.fattening_farm_name_cell,
        state.fattening_farm.name if state.fattening_farm is not None else "",
    )
    set_cell_value(
        ws,
        mp.holding_address_cell,
        f"{holding.address_street} {holding.address_zip} {holding.address_city}",
    )
    set_cell_value(
        ws,
        mp.rearing_farm_name_cell,
        state.rearing_farm.name if state.rearing_farm is not None else "",
    )
    set_cell_value(ws, mp.stall_cell, stalls)
    set_cell_value(ws, mp.parent_flock_cell, format_parent_flocks_by_section(state))
    set_cell_value(
        ws, mp.production_week_cell, format_production_weeks_by_section(state)
    )
    set_cell_value(ws, mp.date_hatched_cell, state.date_hatched)
    set_cell_value(
        ws,
        mp.transfer_date_cell,
        state.transfer.date if state.transfer is not None else "",
    )

    for cell in mp.clear_cells:
        set_cell_value(ws, cell, "")


def opening_time_counts_as_full_outdoor_day(opening_time: str | None) -> bool:
    if opening_time is None:
        return False

    hour_str, separator, minute_str = opening_time.partition(":")
    if separator != ":":
        return False

    try:
        hour = int(hour_str)
        minute = int(minute_str)
    except ValueError:
        return False

    if hour < 0 or hour > 23 or minute < 0 or minute > 59:
        return False

    return (hour, minute) <= (10, 0)


def copy_outdoor_journal_row_style(
    ws: Worksheet, source_row: int, target_row: int
) -> None:
    max_column = ws.max_column if isinstance(ws.max_column, int) else 11
    for col_idx in range(1, max_column + 1):
        source_cell = f"{get_column_letter(col_idx)}{source_row}"
        target_cell = f"{get_column_letter(col_idx)}{target_row}"
        copy_cell_style(ws, source_cell, target_cell)

    try:
        merged_ranges = list(cast(Iterable[CellRange], ws.merged_cells.ranges))
    except AttributeError:
        merged_ranges = []

    row_offset = target_row - source_row
    for merged_range in merged_ranges:
        min_row = cast(int, merged_range.min_row)
        max_row = cast(int, merged_range.max_row)
        if min_row != source_row or max_row != source_row:
            continue

        shifted = CellRange(
            min_col=cast(int, merged_range.min_col),
            min_row=source_row + row_offset,
            max_col=cast(int, merged_range.max_col),
            max_row=source_row + row_offset,
        )
        if any(rng.coord == shifted.coord for rng in merged_ranges):
            continue
        ws.merge_cells(shifted.coord)

    try:
        if ws.row_dimensions[source_row].height is not None:
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    except AttributeError:
        pass


def fill_outdoor_journal_days(
    ws: Worksheet,
    stallkarte: domain.Stallkarte,
    mapping: OutdoorJournalDaysMapping = default_outdoor_journal_days_mapping,
) -> None:
    state = stallkarte.state
    transfer = state.transfer
    if transfer is None:
        set_cell_value(
            ws,
            f"{mapping.summary_total_cell_column}{mapping.summary_row}",
            0,
        )
        return

    fattening_days = [
        day
        for production_day, day in sorted(state.days.items())
        if production_day >= transfer.production_day
    ]

    extra_rows = max(0, len(fattening_days) - mapping.default_data_row_count)
    if extra_rows > 0:
        insertion_start = mapping.first_data_row + mapping.default_data_row_count
        max_column = ws.max_column if isinstance(ws.max_column, int) else 11
        max_row = ws.max_row if isinstance(ws.max_row, int) else insertion_start + 10
        ws.move_range(
            f"A{insertion_start}:{get_column_letter(max_column)}{max_row}",
            rows=extra_rows,
            cols=0,
        )
        shift_merged_ranges_down(ws, insertion_start, extra_rows)

        source_row = insertion_start - 1
        for row in range(insertion_start, insertion_start + extra_rows):
            copy_outdoor_journal_row_style(ws, source_row, row)

    full_outdoor_days_count = 0
    template_last_data_row = mapping.first_data_row + mapping.default_data_row_count - 1
    for index, day in enumerate(fattening_days):
        row = mapping.first_data_row + index
        date = state.date_started + datetime.timedelta(days=day.production_day)

        if row > template_last_data_row:
            copy_outdoor_journal_row_style(ws, template_last_data_row, row)

        set_cell_value(ws, f"{mapping.date_column}{row}", date)
        set_cell_value(
            ws, f"{mapping.opening_time_column}{row}", day.opening_time or ""
        )
        set_cell_value(
            ws,
            f"{mapping.sun_column}{row}",
            "X" if event_models.WeatherCondition.SUN in day.weather_conditions else "",
        )
        set_cell_value(
            ws,
            f"{mapping.cloudy_column}{row}",
            "X"
            if event_models.WeatherCondition.CLOUDY in day.weather_conditions
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.precipitation_column}{row}",
            "X"
            if event_models.WeatherCondition.PRECIPITATION in day.weather_conditions
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.extreme_wetness_column}{row}",
            "X"
            if event_models.WeatherCondition.EXTREME_WETNESS in day.weather_conditions
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.frost_column}{row}",
            "X"
            if event_models.WeatherCondition.FROST in day.weather_conditions
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.strong_wind_column}{row}",
            "X"
            if event_models.WeatherCondition.STRONG_WIND in day.weather_conditions
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.veterinarian_column}{row}",
            "X" if day.veterinarian else "",
        )

        if opening_time_counts_as_full_outdoor_day(day.opening_time):
            full_outdoor_days_count += 1

    summary_row = mapping.summary_row + extra_rows
    set_cell_value(
        ws,
        f"{mapping.summary_total_cell_column}{summary_row}",
        full_outdoor_days_count,
    )


def format_slaughter_name_and_weight(
    slaughterer_name: str | None,
    final_weight_kg: float | None,
) -> str:
    if not slaughterer_name and final_weight_kg is None:
        return ""
    if not slaughterer_name:
        return f"{final_weight_kg:g} kg"
    if final_weight_kg is None:
        return slaughterer_name
    return f"{slaughterer_name} - {final_weight_kg:g} kg"


def fill_throughput_report_finish_notes(
    ws: Worksheet,
    finish_notes: list[NoteEntry],
) -> None:
    mp = default_throughput_report_finish_notes_mapping

    slaughter_notes = [
        note for note in finish_notes if note.note_type == NoteType.SLAUGHTER
    ][: len(mp.slaughter_rows)]
    catching_notes = [
        note for note in finish_notes if note.note_type == NoteType.CATCHING
    ][: len(mp.catching_catcher_cells)]

    for index, row in enumerate(mp.slaughter_rows):
        note = slaughter_notes[index] if index < len(slaughter_notes) else None
        ws[f"{mp.slaughter_date_column}{row}"] = (
            note.slaughter_date if note is not None else ""
        )
        ws[f"{mp.slaughter_animals_count_column}{row}"] = (
            note.slaughter_animals_count
            if note is not None and note.slaughter_animals_count is not None
            else ""
        )
        ws[f"{mp.slaughter_slaughterer_column}{row}"] = (
            format_slaughter_name_and_weight(
                note.slaughterer_name,
                note.slaughter_final_weight_kg,
            )
            if note is not None
            else ""
        )

    for index, catcher_cell in enumerate(mp.catching_catcher_cells):
        note = catching_notes[index] if index < len(catching_notes) else None
        ws[catcher_cell] = note.catcher_name if note is not None else ""

    for index, time_cell in enumerate(mp.catching_time_cells):
        note = catching_notes[index] if index < len(catching_notes) else None
        ws[time_cell] = note.catching_time if note is not None else ""


def fill_throughput_report_general_notes(
    ws: Worksheet,
    stallkarte: domain.Stallkarte,
    mapping: ThroughputReportGeneralNotesMapping = (
        default_throughput_report_general_notes_mapping
    ),
) -> None:
    vaccination_notes = iter_general_notes_by_type(stallkarte, NoteType.VACCINATION)
    treatment_notes = iter_general_notes_by_type(stallkarte, NoteType.TREATMENT)

    for index, row in enumerate(mapping.vaccination_rows):
        item = vaccination_notes[index] if index < len(vaccination_notes) else None
        production_day, note = item if item is not None else (None, None)

        set_cell_value(
            ws,
            f"{mapping.vaccination_delivery_receipt_column}{row}",
            note.delivery_receipt_number if note is not None else "",
        )
        set_cell_value(
            ws,
            f"{mapping.vaccination_batch_number_column}{row}",
            note.batch_number if note is not None else "",
        )
        set_cell_value(
            ws,
            f"{mapping.vaccination_code_column}{row}",
            note.vaccination_code.value.upper()
            if note is not None and note.vaccination_code is not None
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.vaccination_date_column}{row}",
            format_date(
                stallkarte.state.date_started + datetime.timedelta(days=production_day)
            )
            if production_day is not None
            else "",
        )

    for index, row in enumerate(mapping.treatment_rows):
        item = treatment_notes[index] if index < len(treatment_notes) else None
        production_day, note = item if item is not None else (None, None)

        treatment_start_date = (
            stallkarte.state.date_started + datetime.timedelta(days=production_day)
            if production_day is not None
            else None
        )
        waiting_days = waiting_time_to_days(note) if note is not None else None
        treatment_end_date = (
            treatment_start_date + datetime.timedelta(days=waiting_days)
            if treatment_start_date is not None and waiting_days is not None
            else None
        )

        set_cell_value(
            ws,
            f"{mapping.treatment_delivery_receipt_column}{row}",
            note.delivery_receipt_number if note is not None else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_treatment_code_column}{row}",
            note.treatment_code.value
            if note is not None and note.treatment_code is not None
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_amount_column}{row}",
            format_value_and_unit_if_complete(
                note.treatment_amount_value,
                note.treatment_amount_unit.value
                if note is not None and note.treatment_amount_unit is not None
                else None,
            )
            or ""
            if note is not None
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_batch_number_column}{row}",
            note.batch_number if note is not None else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_start_date_column}{row}",
            format_date(treatment_start_date)
            if treatment_start_date is not None
            else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_end_date_column}{row}",
            format_date(treatment_end_date) if treatment_end_date is not None else "",
        )
        set_cell_value(
            ws,
            f"{mapping.treatment_waiting_time_column}{row}",
            format_value_and_unit_if_complete(
                note.treatment_waiting_time_value,
                note.treatment_waiting_time_unit.value
                if note is not None and note.treatment_waiting_time_unit is not None
                else None,
            )
            or ""
            if note is not None
            else "",
        )


class StallkarteWorksheetExporter:
    def __init__(
        self,
        holding: domain.AgriculturalHolding,
        stallkarte: domain.Stallkarte,
        section_number: int,
    ) -> None:
        self.holding = holding
        self.stallkarte = stallkarte
        self.section_number = section_number

    def fill_holding_information(
        self,
        ws: Worksheet,
    ) -> None:
        mp = default_stallkarte_holding_mapping
        state = self.stallkarte.state

        set_cell_value(ws, mp.holding_name, self.holding.name)
        set_cell_value(
            ws,
            mp.holding_address,
            f"{self.holding.address_street} "
            f"{self.holding.address_zip} {self.holding.address_city}",
        )
        set_cell_value(ws, mp.fattening_cycle, state.fattening_cycle)

        if state.rearing_farm is not None:
            set_cell_value(ws, mp.rearing_farm_name, state.rearing_farm.name)
            set_cell_value(
                ws,
                mp.rearing_farm_vvvo_number,
                state.rearing_farm.vvvo_number,
            )

            if section := get_section_from_farm(
                state.rearing_farm, self.section_number
            ):
                set_cell_value(ws, mp.rearing_barn_number, section.number)

        if state.fattening_farm is not None:
            set_cell_value(ws, mp.fattening_farm_name, state.fattening_farm.name)
            set_cell_value(
                ws,
                mp.fattening_farm_vvvo_number,
                state.fattening_farm.vvvo_number,
            )

            if section := get_section_from_farm(
                state.fattening_farm, self.section_number
            ):
                set_cell_value(ws, mp.fattening_barn_number, section.number)

        set_cell_value(ws, mp.date_hatched, state.date_hatched)
        set_cell_value(
            ws,
            mp.animals_count,
            get_initial_animals_count_for_section(state, self.section_number),
        )

        if state.transfer is not None:
            set_cell_value(ws, mp.date_transferred, state.transfer.date)

        set_cell_value(ws, mp.hatchery_name, state.hatchery)
        set_cell_value(ws, mp.breed, state.breed)
        set_cell_value(
            ws,
            mp.production_week,
            format_production_weeks_for_section(state, self.section_number),
        )
        set_cell_value(ws, mp.eco_control_number, state.eco_control_number)
        set_cell_value(
            ws,
            mp.parent_flock,
            format_parent_flocks_for_section(state, self.section_number),
        )

        if state.is_eu_bio:
            eu_bio_ref = resolve_cell_reference(ws, mp.eu_bio)
            set_cell_value(ws, mp.eu_bio, f"{ws[eu_bio_ref].value} X")

        if state.is_naturland:
            naturland_ref = resolve_cell_reference(ws, mp.naturland)
            set_cell_value(ws, mp.naturland, f"{ws[naturland_ref].value} X")

    def fill_checklist(self, ws: Worksheet) -> None:
        mp = default_stallkarte_checklist_mapping
        checklist = self.stallkarte.state.rearing_checklist
        if checklist is None:
            return

        fill_yes_no_cell(
            ws,
            mp.did_emergency_power_test,
            checklist.alarm_test.did_emergency_power_test,
        )
        fill_yes_no_cell(ws, mp.did_alarm_test, checklist.alarm_test.did_alarm_test)
        fill_yes_no_cell(
            ws, mp.did_dark_period_test, checklist.lighting_program.did_dark_period_test
        )
        fill_yes_no_cell(
            ws,
            mp.had_deviations_due_to_vet,
            checklist.lighting_program.had_divergence_due_to_vet,
        )

        if checklist.stable_disinfected.date is not None:
            set_cell_value(
                ws,
                mp.stable_disinfection_date,
                format_date(checklist.stable_disinfected.date),
            )

        set_cell_value(
            ws,
            mp.stable_disinfectant,
            checklist.stable_disinfected.disinfectant,
        )
        set_cell_value(
            ws,
            mp.stable_disinfectant_dosis,
            checklist.stable_disinfected.dosis,
        )

        if checklist.silo_cleaned.date is not None:
            set_cell_value(
                ws,
                mp.silo_cleaning_date,
                format_date(checklist.silo_cleaned.date),
            )

        set_cell_value(ws, mp.silo_detergent, checklist.silo_cleaned.detergent)
        set_cell_value(ws, mp.silo_detergent_dosis, checklist.silo_cleaned.dosis)

        if checklist.water_line_disinfected.date is not None:
            set_cell_value(
                ws,
                mp.water_line_disinfection_date,
                format_date(checklist.water_line_disinfected.date),
            )

        set_cell_value(
            ws,
            mp.water_line_disinfectant,
            checklist.water_line_disinfected.disinfectant,
        )
        set_cell_value(
            ws,
            mp.water_line_disinfectant_dosis,
            checklist.water_line_disinfected.dosis,
        )

        fill_yes_no_cell(
            ws,
            mp.did_pest_control_measures,
            checklist.pest_control_measures.did_perform_pest_control,
        )
        set_cell_value(
            ws,
            mp.pest_control_measures_annotation,
            checklist.pest_control_measures.annotation,
        )

    def fill_production_day(self, ws: Worksheet, day: StallkarteStateDay) -> None:
        mp = get_mapping_of_production_day(day.production_day)

        section = day.sections.get(self.section_number)

        set_cell_value(ws, mp.day, day.production_day)
        set_cell_value(ws, mp.temperature, day.temperature_celsius)
        if day.humidity_percent is not None:
            set_cell_value(ws, mp.humidity, day.humidity_percent / 100)
        set_cell_value(ws, mp.water_consumption, day.water_consumption_liters)
        set_cell_value(ws, mp.feed_consumption, day.feed_consumption_kg)
        set_cell_value(ws, mp.weight, day.weight_grams)

        formatted_notes = [format_general_note_entry(note) for note in day.notes]
        note_cell_value = " / ".join(formatted_notes)

        if section is not None:
            set_cell_value(
                ws,
                mp.natural_mortality,
                optional_sum(
                    section.natural_mortality_morning,
                    section.natural_mortality_evening,
                ),
            )
            set_cell_value(
                ws,
                mp.selective_mortality,
                optional_sum(
                    section.selective_mortality_morning,
                    section.selective_mortality_evening,
                ),
            )
            set_cell_value(
                ws,
                mp.did_inspection_1,
                "X" if section.did_inspection_morning else "",
            )
            set_cell_value(
                ws,
                mp.did_inspection_2,
                "X" if section.did_inspection_evening else "",
            )

        set_cell_value(ws, mp.annotation, note_cell_value)

    def fill_production_days(self, ws: Worksheet) -> None:
        for day_number, day in self.stallkarte.state.days.items():
            self.fill_production_day(ws, day)


class Exporter:
    def __init__(
        self, template_file: str, template_file_default_worksheet: str
    ) -> None:
        if not template_file:
            raise ValueError("template file path must be provided.")
        if not template_file_default_worksheet:
            raise ValueError("template file worksheet must be provided.")

        self.template_file: str = template_file
        self.template_file_default_worksheet: str = template_file_default_worksheet

    def export(
        self,
        holding: domain.AgriculturalHolding,
        stallkarte: domain.Stallkarte,
    ) -> Workbook:
        wb: Workbook = openpyxl.load_workbook(self.template_file)
        ws_template = wb[self.template_file_default_worksheet]

        if stallkarte.state.rearing_farm is None:
            raise ValueError("stallkarte state must have a rearing farm to export.")

        if stallkarte.state.fattening_farm is not None:
            if len(stallkarte.state.rearing_farm.sections) != len(
                stallkarte.state.fattening_farm.sections
            ):
                raise ValueError(
                    "the number of sections in the rearing farm and the fattening farm "
                    "must be the same."
                )

        pairs: list[tuple[StallkarteWorksheetExporter, Worksheet]] = []

        # To enable export when the fattening farm is not yet assigned, we use
        # the rearing farm only
        sections = stallkarte.state.rearing_farm.sections

        # Insert the new worksheets at the same position of the template
        right_offset = len(wb.worksheets) - wb.worksheets.index(ws_template) - 1

        for section in sections:
            ws = wb.copy_worksheet(ws_template)
            ws.title = (
                f"{default_export_worksheet_names.section_title_prefix} "
                f"{section.number}"
            )
            wb.move_sheet(ws, offset=-right_offset)

            pairs.append(
                (
                    StallkarteWorksheetExporter(holding, stallkarte, section.number),
                    ws,
                )
            )

        wb.remove(ws_template)

        for exporter, ws in pairs:
            exporter.fill_holding_information(
                ws,
            )
            exporter.fill_checklist(ws)
            exporter.fill_production_days(ws)

        section_sheet_names = [worksheet.title for _, worksheet in pairs]

        quality_report_name = (
            default_export_worksheet_names.quality_report_bio_chicks_name
        )
        if quality_report_name in wb.sheetnames:
            fill_quality_report_bio_chicks(
                wb[quality_report_name],
                holding,
                stallkarte,
                section_sheet_names,
            )

        outdoor_journal_name = default_export_worksheet_names.outdoor_journal_name
        if outdoor_journal_name in wb.sheetnames:
            ws_outdoor_journal = wb[outdoor_journal_name]
            fill_outdoor_journal_header(ws_outdoor_journal, holding, stallkarte)
            fill_outdoor_journal_days(ws_outdoor_journal, stallkarte)

        ws_throughput_report = wb[default_export_worksheet_names.throughput_report_name]
        fill_throughput_report_header(ws_throughput_report, holding, stallkarte)
        fill_throughput_report_total_animals(
            ws_throughput_report,
            section_sheet_names,
        )
        fill_throughput_report_mortality(
            ws_throughput_report,
            section_sheet_names,
        )
        fill_throughput_report_general_notes(
            ws_throughput_report,
            stallkarte,
        )
        fill_throughput_report_finish_notes(
            ws_throughput_report,
            stallkarte.state.finish_notes,
        )

        return wb
