"""
This is not great and generate by Gemini, but it is better than nothing.
And sheet export is a bit hard to test without actually writing to a file,
which we want to avoid in unit tests.
"""

import datetime
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from app import domain
from app.core.stallkarteexport.exporter import (
    StallkarteWorksheetExporter,
    build_sheet_sum_formula,
    copy_outdoor_journal_row_style,
    fill_outdoor_journal_days,
    fill_quality_report_bio_chicks,
    fill_outdoor_journal_header,
    fill_throughput_report_general_notes,
    fill_throughput_report_finish_notes,
    fill_throughput_report_header,
    fill_throughput_report_mortality,
    fill_throughput_report_total_animals,
    fill_yes_no_cell,
    format_slaughter_name_and_weight,
    format_date,
    optional_sum,
    shift_merged_ranges_down,
)
from app.core.stallkarteexport.mapping import (
    YesNoMapping,
    default_stallkarte_checklist_mapping,
    default_stallkarte_holding_mapping,
    get_mapping_of_production_day,
    get_throughput_report_mortality_row_pairs,
)
from app.domain.stallkarte.events.models import (
    NoteEntry,
    NoteType,
    ParentFlockEntry,
    SectionInstallationDetails,
    SockTestResult,
    TreatmentAmountUnit,
    TreatmentCode,
    VaccinationCode,
    WeatherCondition,
    WaitingTimeUnit,
)
from app.domain import StallkarteStateDay


def test_format_date() -> None:
    date = datetime.date(2023, 10, 27)
    assert format_date(date) == "27.10.2023"


def test_optional_sum() -> None:
    assert optional_sum(None, None) is None
    assert optional_sum(1, None) == 1
    assert optional_sum(None, 2) == 2
    assert optional_sum(1, 2) == 3


def test_fill_yes_no_cell() -> None:
    ws = MagicMock(spec=Worksheet)
    mapping = YesNoMapping(yes="A1", no="B1")

    # Test True
    fill_yes_no_cell(ws, mapping, True)
    ws.__setitem__.assert_any_call("A1", "Ja")
    assert ws["A1"].font == Font(bold=True)
    ws.__setitem__.assert_any_call("B1", "")

    ws.reset_mock()

    # Test False
    fill_yes_no_cell(ws, mapping, False)
    ws.__setitem__.assert_any_call("A1", "")
    ws.__setitem__.assert_any_call("B1", "Nein")
    assert ws["B1"].font == Font(bold=True)

    ws.reset_mock()

    # Test None (treated as False in current implementation)
    fill_yes_no_cell(ws, mapping, None)
    ws.__setitem__.assert_any_call("A1", "")
    ws.__setitem__.assert_any_call("B1", "Nein")
    assert ws["B1"].font == Font(bold=True)


def test_shift_merged_ranges_down() -> None:
    wb = Workbook()
    ws = wb.active

    ws.merge_cells("A30:E30")
    ws.merge_cells("A33:J38")
    ws.merge_cells("A41:J41")
    ws.merge_cells("B42:J42")

    shift_merged_ranges_down(ws, start_row=31, offset=1)

    merged_coords = {cell_range.coord for cell_range in ws.merged_cells.ranges}
    assert "A30:E30" in merged_coords
    assert "A34:J39" in merged_coords
    assert "A42:J42" in merged_coords
    assert "B43:J43" in merged_coords


def test_copy_outdoor_journal_row_style_preserves_row_merges() -> None:
    wb = Workbook()
    ws = wb.active

    ws.merge_cells("C10:D10")
    ws["C10"] = "merged"

    copy_outdoor_journal_row_style(ws, source_row=10, target_row=11)

    merged_coords = {cell_range.coord for cell_range in ws.merged_cells.ranges}
    assert "C11:D11" in merged_coords
    assert "D11:E11" not in merged_coords


def test_build_sheet_sum_formula() -> None:
    assert build_sheet_sum_formula([], "E15") == "=0"
    assert (
        build_sheet_sum_formula(["Stallkarte 1"], "E15") == "=SUM('Stallkarte 1'!E15)"
    )
    assert (
        build_sheet_sum_formula(["Stallkarte 1", "Stallkarte 2"], "E15")
        == "=SUM('Stallkarte 1'!E15,'Stallkarte 2'!E15)"
    )


def test_get_throughput_report_mortality_row_pairs() -> None:
    rearing_pairs, fattening_pairs = get_throughput_report_mortality_row_pairs()

    assert len(rearing_pairs) == len(fattening_pairs) == 35
    assert rearing_pairs[0] == (12, 15)
    assert rearing_pairs[-1] == (46, 50)
    assert fattening_pairs[0] == (12, 51)
    assert fattening_pairs[-1] == (46, 87)


def test_fill_throughput_report_mortality() -> None:
    ws = MagicMock(spec=Worksheet)

    fill_throughput_report_mortality(ws, ["Stallkarte 1", "Stallkarte 2"])

    ws.__setitem__.assert_any_call("I12", "=SUM('Stallkarte 1'!E15,'Stallkarte 2'!E15)")
    ws.__setitem__.assert_any_call("I20", "=SUM('Stallkarte 1'!E24,'Stallkarte 2'!E24)")
    ws.__setitem__.assert_any_call("I46", "=SUM('Stallkarte 1'!E50,'Stallkarte 2'!E50)")

    ws.__setitem__.assert_any_call("L12", "=SUM('Stallkarte 1'!E51,'Stallkarte 2'!E51)")
    ws.__setitem__.assert_any_call("L17", "=SUM('Stallkarte 1'!E57,'Stallkarte 2'!E57)")
    ws.__setitem__.assert_any_call("L24", "=SUM('Stallkarte 1'!E65,'Stallkarte 2'!E65)")
    ws.__setitem__.assert_any_call("L46", "=SUM('Stallkarte 1'!E87,'Stallkarte 2'!E87)")

    # Ensure only the expected report windows are touched.
    touched_cells = {call.args[0] for call in ws.__setitem__.call_args_list}
    assert "I11" not in touched_cells
    assert "I47" not in touched_cells
    assert "L11" not in touched_cells
    assert "L47" not in touched_cells


def test_fill_throughput_report_total_animals() -> None:
    ws = MagicMock(spec=Worksheet)

    fill_throughput_report_total_animals(ws, ["Stallkarte 1", "Stallkarte 2"])

    ws.__setitem__.assert_any_call("D7", "=SUM('Stallkarte 1'!I4,'Stallkarte 2'!I4)")


def test_fill_throughput_report_header() -> None:
    ws = MagicMock(spec=Worksheet)

    holding = MagicMock(spec=domain.AgriculturalHolding)
    holding.name = "Test Holding"
    holding.address_street = "Street"
    holding.address_zip = "12345"
    holding.address_city = "City"

    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.fattening_cycle = "2026-A"
    state.fattening_farm = MagicMock(spec=domain.StallkarteFarm)
    state.fattening_farm.name = "Mast Farm"
    state.fattening_farm.vvvo_number = "VVVO-F"
    state.fattening_farm.sections = [MagicMock(number=1), MagicMock(number=2)]
    state.rearing_farm = MagicMock(spec=domain.StallkarteFarm)
    state.rearing_farm.name = "Aufzucht Farm"
    state.rearing_farm.vvvo_number = "VVVO-R"
    state.rearing_farm.sections = [MagicMock(number=1), MagicMock(number=2)]
    state.date_hatched = datetime.date(2026, 3, 1)
    state.transfer = MagicMock()
    state.transfer.date = datetime.date(2026, 3, 20)
    state.hatchery = "Seed Hatchery"
    state.breed = "Ross 308"
    state.eco_control_number = "DE-OKO-001"
    state.installation_details_by_section = {
        1: SectionInstallationDetails(
            section_number=1,
            initial_animals_count=1000,
            initial_weight_grams=42.5,
            bedding="Strohhäcksel",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="et12", production_week=12),
                ParentFlockEntry(herd_identifier="et13", production_week=13),
            ],
        ),
        2: SectionInstallationDetails(
            section_number=2,
            initial_animals_count=980,
            initial_weight_grams=43.0,
            bedding="Strohhäcksel",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="et22", production_week=22)
            ],
        ),
    }
    state.is_eu_bio = True
    state.is_naturland = False

    fill_throughput_report_header(ws, holding, stallkarte)

    ws.__setitem__.assert_any_call("B4", "Test Holding")
    ws.__setitem__.assert_any_call("D4", "Street 12345 City")
    ws.__setitem__.assert_any_call("L4", "2026-A")
    ws.__setitem__.assert_any_call("B5", "Mast Farm")
    ws.__setitem__.assert_any_call("D5", "1, 2")
    ws.__setitem__.assert_any_call("H5", "VVVO-F")
    ws.__setitem__.assert_any_call("B6", "Aufzucht Farm")
    ws.__setitem__.assert_any_call("D6", "1, 2")
    ws.__setitem__.assert_any_call("H6", "VVVO-R")
    ws.__setitem__.assert_any_call("B7", datetime.date(2026, 3, 1))
    ws.__setitem__.assert_any_call("H7", datetime.date(2026, 3, 20))
    ws.__setitem__.assert_any_call("B8", "Seed Hatchery")
    ws.__setitem__.assert_any_call("D8", "Ross 308")
    ws.__setitem__.assert_any_call("H8", "12; 13")
    ws.__setitem__.assert_any_call("J8", "22")
    ws.__setitem__.assert_any_call("L8", "")
    ws.__setitem__.assert_any_call("B9", "DE-OKO-001")
    ws.__setitem__.assert_any_call("D9", "et12; et13")
    ws.__setitem__.assert_any_call("E9", "et22")
    ws.__setitem__.assert_any_call("F9", "")
    ws.__setitem__.assert_any_call("B41", "Strohhäcksel")
    ws.__setitem__.assert_any_call("H9", "X")
    ws.__setitem__.assert_any_call("L9", "")


def test_fill_quality_report_bio_chicks() -> None:
    ws = MagicMock(spec=Worksheet)

    holding = MagicMock(spec=domain.AgriculturalHolding)
    holding.name = "Test Holding"
    holding.address_street = "Street"
    holding.address_zip = "12345"
    holding.address_city = "City"

    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.rearing_farm = MagicMock(spec=domain.StallkarteFarm)
    state.rearing_farm.name = "Rearing Farm"
    section_1 = MagicMock()
    section_1.number = 1
    section_1.name = "Abteil 1"
    section_2 = MagicMock()
    section_2.number = 2
    section_2.name = "Abteil 2"
    state.rearing_farm.sections = [section_1, section_2]
    state.fattening_farm = None
    state.installation_details_by_section = {
        1: SectionInstallationDetails(
            section_number=1,
            initial_animals_count=1000,
            initial_weight_grams=42.5,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-1", production_week=12),
                ParentFlockEntry(herd_identifier="PF-2", production_week=13),
            ],
        ),
        2: SectionInstallationDetails(
            section_number=2,
            initial_animals_count=970,
            initial_weight_grams=43.0,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-3", production_week=14)
            ],
        ),
    }
    state.date_hatched = datetime.date(2026, 1, 10)
    state.hatchery = "Hatchery X"

    fill_quality_report_bio_chicks(
        ws, holding, stallkarte, ["Stallkarte 1", "Stallkarte 2"]
    )

    ws.__setitem__.assert_any_call("B3", "Test Holding")
    ws.__setitem__.assert_any_call("B5", "Street 12345 City")
    ws.__setitem__.assert_any_call("B7", "Rearing Farm")
    ws.__setitem__.assert_any_call("B8", "1")
    ws.__setitem__.assert_any_call("E8", "2")
    ws.__setitem__.assert_any_call("B17", "Stall 1")
    ws.__setitem__.assert_any_call("E17", "Stall 2")
    ws.__setitem__.assert_any_call("H17", "")
    ws.__setitem__.assert_any_call("B9", "PF-1; PF-2")
    ws.__setitem__.assert_any_call("E9", "PF-3")
    ws.__setitem__.assert_any_call("B10", "12; 13")
    ws.__setitem__.assert_any_call("E10", "14")
    ws.__setitem__.assert_any_call("C15", "=SUM('Stallkarte 1'!I4,'Stallkarte 2'!I4)")

    ws.__setitem__.assert_any_call("B19", "='Stallkarte 1'!C15")
    ws.__setitem__.assert_any_call("C19", "='Stallkarte 1'!D15")
    ws.__setitem__.assert_any_call("D19", "='Stallkarte 1'!E15")

    ws.__setitem__.assert_any_call("E19", "='Stallkarte 2'!C15")
    ws.__setitem__.assert_any_call("F19", "='Stallkarte 2'!D15")
    ws.__setitem__.assert_any_call("G19", "='Stallkarte 2'!E15")
    ws.move_range.assert_called_once_with("A31:K60", rows=1, cols=0)
    ws.__setitem__.assert_any_call(
        "A30", "Durchschnittsgewicht am 7. Tag in Gramm - Abteil 1"
    )
    ws.__setitem__.assert_any_call(
        "A31", "Durchschnittsgewicht am 7. Tag in Gramm - Abteil 2"
    )
    ws.__setitem__.assert_any_call("F30", "='Stallkarte 1'!H22")
    ws.__setitem__.assert_any_call("F31", "='Stallkarte 2'!H22")

    ws.__setitem__.assert_any_call("H19", "")
    ws.__setitem__.assert_any_call("I19", "")
    ws.__setitem__.assert_any_call("J19", "")


def test_fill_quality_report_bio_chicks_with_four_sections() -> None:
    ws = MagicMock(spec=Worksheet)

    holding = MagicMock(spec=domain.AgriculturalHolding)
    holding.name = "Test Holding"
    holding.address_street = "Street"
    holding.address_zip = "12345"
    holding.address_city = "City"

    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.rearing_farm = MagicMock(spec=domain.StallkarteFarm)
    state.rearing_farm.name = "Rearing Farm"
    section_1 = MagicMock()
    section_1.number = 1
    section_1.name = "Abteil 1"
    section_2 = MagicMock()
    section_2.number = 2
    section_2.name = "Abteil 2"
    section_3 = MagicMock()
    section_3.number = 3
    section_3.name = "Abteil 3"
    section_4 = MagicMock()
    section_4.number = 4
    section_4.name = "Abteil 4"
    state.rearing_farm.sections = [section_1, section_2, section_3, section_4]
    state.fattening_farm = None
    state.installation_details_by_section = {
        1: SectionInstallationDetails(
            section_number=1,
            initial_animals_count=1000,
            initial_weight_grams=42.5,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-1", production_week=12)
            ],
        ),
        2: SectionInstallationDetails(
            section_number=2,
            initial_animals_count=980,
            initial_weight_grams=43.0,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-2", production_week=13)
            ],
        ),
        3: SectionInstallationDetails(
            section_number=3,
            initial_animals_count=960,
            initial_weight_grams=44.0,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-3", production_week=14)
            ],
        ),
        4: SectionInstallationDetails(
            section_number=4,
            initial_animals_count=940,
            initial_weight_grams=45.0,
            bedding="Stroh",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="PF-4", production_week=15)
            ],
        ),
    }
    state.date_hatched = datetime.date(2026, 1, 10)
    state.hatchery = "Hatchery X"

    fill_quality_report_bio_chicks(
        ws,
        holding,
        stallkarte,
        ["Stallkarte 1", "Stallkarte 2", "Stallkarte 3", "Stallkarte 4"],
    )

    ws.__setitem__.assert_any_call("K17", "Stall 4")
    ws.__setitem__.assert_any_call("K19", "='Stallkarte 4'!C15")
    ws.__setitem__.assert_any_call("L19", "='Stallkarte 4'!D15")
    ws.__setitem__.assert_any_call("M19", "='Stallkarte 4'!E15")
    ws.__setitem__.assert_any_call("K9", "PF-4")
    ws.__setitem__.assert_any_call("K10", "15")
    ws.move_range.assert_called_once_with("A31:K60", rows=3, cols=0)
    ws.__setitem__.assert_any_call("F30", "='Stallkarte 1'!H22")
    ws.__setitem__.assert_any_call("F31", "='Stallkarte 2'!H22")
    ws.__setitem__.assert_any_call("F32", "='Stallkarte 3'!H22")
    ws.__setitem__.assert_any_call("F33", "='Stallkarte 4'!H22")


def test_fill_outdoor_journal_header() -> None:
    ws = MagicMock(spec=Worksheet)

    holding = MagicMock(spec=domain.AgriculturalHolding)
    holding.name = "Test Holding"
    holding.address_street = "Street"
    holding.address_zip = "12345"
    holding.address_city = "City"

    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.fattening_cycle = "2026-A"
    state.rearing_farm = MagicMock(spec=domain.StallkarteFarm)
    state.rearing_farm.name = "Aufzucht Farm"
    state.rearing_farm.sections = [MagicMock(number=1), MagicMock(number=2)]
    state.fattening_farm = MagicMock(spec=domain.StallkarteFarm)
    state.fattening_farm.name = "Mast Farm"
    state.installation_details_by_section = {
        1: SectionInstallationDetails(
            section_number=1,
            initial_animals_count=1000,
            initial_weight_grams=42.5,
            bedding="Strohhäcksel",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="et12", production_week=12),
                ParentFlockEntry(herd_identifier="et13", production_week=13),
            ],
        ),
        2: SectionInstallationDetails(
            section_number=2,
            initial_animals_count=980,
            initial_weight_grams=43.0,
            bedding="Strohhäcksel",
            parent_flocks=[
                ParentFlockEntry(herd_identifier="et22", production_week=22)
            ],
        ),
    }
    state.date_hatched = datetime.date(2026, 3, 1)
    state.transfer = MagicMock()
    state.transfer.date = datetime.date(2026, 3, 20)

    fill_outdoor_journal_header(ws, holding, stallkarte)

    ws.__setitem__.assert_any_call("K3", "2026-A")
    ws.__setitem__.assert_any_call("B4", "Test Holding")
    ws.__setitem__.assert_any_call("I4", "Mast Farm")
    ws.__setitem__.assert_any_call("B5", "Street 12345 City")
    ws.__setitem__.assert_any_call("I5", "Aufzucht Farm")
    ws.__setitem__.assert_any_call("A7", "1, 2")
    ws.__setitem__.assert_any_call("B7", "1: et12; et13 | 2: et22")
    ws.__setitem__.assert_any_call("F7", "1: 12; 13 | 2: 22")
    ws.__setitem__.assert_any_call("I7", datetime.date(2026, 3, 1))
    ws.__setitem__.assert_any_call("K7", datetime.date(2026, 3, 20))


def test_fill_outdoor_journal_days() -> None:
    ws = MagicMock(spec=Worksheet)
    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.date_started = datetime.date(2026, 3, 1)
    state.transfer = MagicMock(spec=domain.StallkarteTransfer)
    state.transfer.production_day = 2
    state.days = {
        1: MagicMock(
            production_day=1,
            opening_time="08:00",
            weather_conditions=[WeatherCondition.SUN],
            veterinarian=True,
        ),
        2: MagicMock(
            production_day=2,
            opening_time="09:45",
            weather_conditions=[WeatherCondition.SUN, WeatherCondition.EXTREME_WETNESS],
            veterinarian=True,
        ),
        3: MagicMock(
            production_day=3,
            opening_time="10:30",
            weather_conditions=[WeatherCondition.CLOUDY, WeatherCondition.STRONG_WIND],
            veterinarian=False,
        ),
    }

    fill_outdoor_journal_days(ws, stallkarte)

    ws.__setitem__.assert_any_call("A10", datetime.date(2026, 3, 3))
    ws.__setitem__.assert_any_call("B10", "09:45")
    ws.__setitem__.assert_any_call("C10", "X")
    ws.__setitem__.assert_any_call("G10", "X")
    ws.__setitem__.assert_any_call("J10", "X")

    ws.__setitem__.assert_any_call("A11", datetime.date(2026, 3, 4))
    ws.__setitem__.assert_any_call("E11", "X")
    ws.__setitem__.assert_any_call("I11", "X")
    ws.__setitem__.assert_any_call("J11", "")

    ws.__setitem__.assert_any_call("G13", 1)


def test_fill_outdoor_journal_days_expands_rows_when_needed() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Auslaufjournal"
    ws["A12"] = "Summe"
    ws["K19"] = "footer"

    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.date_started = datetime.date(2026, 1, 1)
    state.transfer = MagicMock(spec=domain.StallkarteTransfer)
    state.transfer.production_day = 0

    days: dict[int, MagicMock] = {}
    for production_day in range(35):
        days[production_day] = MagicMock(
            production_day=production_day,
            opening_time="09:30" if production_day % 2 == 0 else "10:30",
            weather_conditions=[],
            veterinarian=False,
        )
    state.days = days

    fill_outdoor_journal_days(ws, stallkarte)

    assert ws["A46"].value == "Summe"
    assert ws["A44"].value == datetime.date(2026, 2, 4)
    assert ws["G46"].value == 18
    assert ws["K53"].value == "footer"


def test_format_slaughter_name_and_weight() -> None:
    assert format_slaughter_name_and_weight(None, None) == ""
    assert format_slaughter_name_and_weight("Steinfelder", None) == "Steinfelder"
    assert format_slaughter_name_and_weight(None, 2.5) == "2.5 kg"
    assert (
        format_slaughter_name_and_weight("Steinfelder", 2.5) == "Steinfelder - 2.5 kg"
    )


def test_fill_throughput_report_finish_notes() -> None:
    ws = MagicMock(spec=Worksheet)
    finish_notes = [
        NoteEntry(
            id="sl-1",
            note_type=NoteType.SLAUGHTER,
            slaughter_date=datetime.date(2026, 3, 20),
            slaughter_animals_count=1200,
            slaughter_final_weight_kg=2.4,
            slaughterer_name="Steinfelder",
        ),
        NoteEntry(
            id="ca-1",
            note_type=NoteType.CATCHING,
            catching_time="05:45",
            catcher_name="Team Nord",
        ),
    ]

    fill_throughput_report_finish_notes(ws, finish_notes)

    ws.__setitem__.assert_any_call("B11", datetime.date(2026, 3, 20))
    ws.__setitem__.assert_any_call("C11", 1200)
    ws.__setitem__.assert_any_call("E11", "Steinfelder - 2.4 kg")
    ws.__setitem__.assert_any_call("B12", "")
    ws.__setitem__.assert_any_call("C12", "")
    ws.__setitem__.assert_any_call("E12", "")
    ws.__setitem__.assert_any_call("B13", "")
    ws.__setitem__.assert_any_call("C13", "")
    ws.__setitem__.assert_any_call("E13", "")
    ws.__setitem__.assert_any_call("B38", "Team Nord")
    ws.__setitem__.assert_any_call("B39", "05:45")
    ws.__setitem__.assert_any_call("C38", "")
    ws.__setitem__.assert_any_call("C39", "")


def test_fill_throughput_report_general_notes() -> None:
    ws = MagicMock(spec=Worksheet)
    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.date_started = datetime.date(2026, 3, 1)
    state.days = {
        2: MagicMock(
            notes=[
                NoteEntry(
                    id="vacc-1",
                    note_type=NoteType.VACCINATION,
                    delivery_receipt_number="AB-11",
                    batch_number="CH-11",
                    vaccination_code=VaccinationCode.ND,
                )
            ]
        ),
        4: MagicMock(
            notes=[
                NoteEntry(
                    id="treat-1",
                    note_type=NoteType.TREATMENT,
                    delivery_receipt_number="AB-22",
                    treatment_code=TreatmentCode.AMPROLINE,
                    treatment_amount_value=1.5,
                    treatment_amount_unit=TreatmentAmountUnit.L_PER_1000,
                    batch_number="CH-22",
                    treatment_waiting_time_value=2,
                    treatment_waiting_time_unit=WaitingTimeUnit.WEEK,
                )
            ]
        ),
    }

    fill_throughput_report_general_notes(ws, stallkarte)

    ws.__setitem__.assert_any_call("A21", "AB-11")
    ws.__setitem__.assert_any_call("B21", "CH-11")
    ws.__setitem__.assert_any_call("C21", "ND")
    ws.__setitem__.assert_any_call("G21", "03.03.2026")

    ws.__setitem__.assert_any_call("A30", "AB-22")
    ws.__setitem__.assert_any_call("B30", "amproline")
    ws.__setitem__.assert_any_call("C30", "1.5 l/1000")
    ws.__setitem__.assert_any_call("D30", "CH-22")
    ws.__setitem__.assert_any_call("E30", "05.03.2026")
    ws.__setitem__.assert_any_call("F30", "19.03.2026")
    ws.__setitem__.assert_any_call("G30", "2 week")


def test_fill_throughput_general_notes_waiting_fields_blank() -> None:
    ws = MagicMock(spec=Worksheet)
    stallkarte = MagicMock(spec=domain.Stallkarte)
    state = MagicMock(spec=domain.StallkarteState)
    stallkarte.state = state
    state.date_started = datetime.date(2026, 3, 1)
    state.days = {
        1: MagicMock(
            notes=[
                NoteEntry(
                    id="treat-1",
                    note_type=NoteType.TREATMENT,
                    treatment_code=TreatmentCode.AMPROLINE,
                )
            ]
        )
    }

    fill_throughput_report_general_notes(ws, stallkarte)

    ws.__setitem__.assert_any_call("E30", "02.03.2026")
    ws.__setitem__.assert_any_call("F30", "")
    ws.__setitem__.assert_any_call("G30", "")


class TestStallkarteWorksheetExporter:
    @pytest.fixture
    def mock_holding(self) -> MagicMock:
        holding = MagicMock(spec=domain.AgriculturalHolding)
        holding.name = "Test Holding"
        holding.address_street = "Test Street"
        holding.address_zip = "12345"
        holding.address_city = "Test City"
        return holding

    @pytest.fixture
    def mock_stallkarte(self) -> MagicMock:
        stallkarte = MagicMock(spec=domain.Stallkarte)
        state = MagicMock(spec=domain.StallkarteState)
        stallkarte.state = state

        # Setup default state values
        state.fattening_cycle = "Cycle 1"
        state.rearing_farm = MagicMock(spec=domain.StallkarteFarm)
        state.rearing_farm.name = "Rearing Farm"
        state.rearing_farm.vvvo_number = "123"
        rearing_section_1 = MagicMock(number=1)
        rearing_section_1.name = "Abschnitt 1"
        rearing_section_2 = MagicMock(number=2)
        rearing_section_2.name = "Abschnitt 2"
        state.rearing_farm.sections = [rearing_section_1, rearing_section_2]

        state.fattening_farm = MagicMock(spec=domain.StallkarteFarm)
        state.fattening_farm.name = "Fattening Farm"
        state.fattening_farm.vvvo_number = "456"
        fattening_section_1 = MagicMock(number=1)
        fattening_section_1.name = "Abschnitt 1"
        fattening_section_2 = MagicMock(number=2)
        fattening_section_2.name = "Abschnitt 2"
        state.fattening_farm.sections = [fattening_section_1, fattening_section_2]

        state.date_hatched = datetime.date(2023, 1, 1)
        state.transfer = MagicMock()
        state.transfer.date = datetime.date(2023, 2, 1)
        state.hatchery = "Test Hatchery"
        state.breed = "Test Breed"
        state.eco_control_number = "ECO123"
        state.installation_details_by_section = {
            1: SectionInstallationDetails(
                section_number=1,
                initial_animals_count=1111,
                initial_weight_grams=42.5,
                bedding="Strohpellets",
                parent_flocks=[
                    ParentFlockEntry(
                        herd_identifier="Parent Flock A", production_week=1
                    ),
                    ParentFlockEntry(
                        herd_identifier="Parent Flock B", production_week=2
                    ),
                ],
            ),
            2: SectionInstallationDetails(
                section_number=2,
                initial_animals_count=999,
                initial_weight_grams=43.0,
                bedding="Strohpellets",
                parent_flocks=[
                    ParentFlockEntry(
                        herd_identifier="Parent Flock C", production_week=3
                    )
                ],
            ),
        }
        state.is_eu_bio = True
        state.is_naturland = False
        state.finish_notes = []

        state.rearing_checklist = MagicMock()
        state.rearing_checklist.alarm_test.did_emergency_power_test = True
        state.rearing_checklist.alarm_test.did_alarm_test = False
        state.rearing_checklist.lighting_program.did_dark_period_test = True
        state.rearing_checklist.lighting_program.had_divergence_due_to_vet = False
        state.rearing_checklist.stable_disinfected.date = datetime.date(2023, 1, 10)
        state.rearing_checklist.stable_disinfected.disinfectant = "Disinfectant A"
        state.rearing_checklist.stable_disinfected.dosis = "10ml"
        state.rearing_checklist.silo_cleaned.date = datetime.date(2023, 1, 11)
        state.rearing_checklist.silo_cleaned.detergent = "Detergent B"
        state.rearing_checklist.silo_cleaned.dosis = "20ml"
        state.rearing_checklist.water_line_disinfected.date = datetime.date(2023, 1, 12)
        state.rearing_checklist.water_line_disinfected.disinfectant = "Disinfectant C"
        state.rearing_checklist.water_line_disinfected.dosis = "30ml"
        state.rearing_checklist.pest_control_measures.did_perform_pest_control = True
        state.rearing_checklist.pest_control_measures.annotation = "Pest control note"

        state.days = {}

        return stallkarte

    def test_fill_holding_information(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)
        # Mock cell access to return a MagicMock that has a value attribute
        cell_mock = MagicMock()
        cell_mock.value = "Original"
        ws.__getitem__.return_value = cell_mock

        exporter.fill_holding_information(ws)

        mp = default_stallkarte_holding_mapping

        # Verify assignments
        ws.__setitem__.assert_any_call(mp.holding_name, "Test Holding")
        ws.__setitem__.assert_any_call(
            mp.holding_address, "Test Street 12345 Test City"
        )
        ws.__setitem__.assert_any_call(mp.fattening_cycle, "Cycle 1")
        ws.__setitem__.assert_any_call(mp.rearing_farm_name, "Rearing Farm")
        ws.__setitem__.assert_any_call(mp.rearing_farm_vvvo_number, "123")
        ws.__setitem__.assert_any_call(mp.rearing_barn_number, 1)
        ws.__setitem__.assert_any_call(mp.fattening_farm_name, "Fattening Farm")
        ws.__setitem__.assert_any_call(mp.fattening_farm_vvvo_number, "456")
        ws.__setitem__.assert_any_call(mp.fattening_barn_number, 1)
        ws.__setitem__.assert_any_call(mp.date_hatched, datetime.date(2023, 1, 1))
        ws.__setitem__.assert_any_call(mp.animals_count, 1111)
        ws.__setitem__.assert_any_call(mp.date_transferred, datetime.date(2023, 2, 1))
        ws.__setitem__.assert_any_call(mp.hatchery_name, "Test Hatchery")
        ws.__setitem__.assert_any_call(mp.breed, "Test Breed")
        ws.__setitem__.assert_any_call(mp.production_week, "1; 2")
        ws.__setitem__.assert_any_call(mp.eco_control_number, "ECO123")
        ws.__setitem__.assert_any_call(
            mp.parent_flock, "Parent Flock A; Parent Flock B"
        )

        # Check eu_bio append
        ws.__setitem__.assert_any_call(mp.eu_bio, "Original X")

    def test_fill_holding_information_uses_transfer_animals_when_installation_missing(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        mock_stallkarte.state.installation_details_by_section = {}
        mock_stallkarte.state.transfer.animals_by_section_number = {1: 1234, 2: 1100}

        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)
        cell_mock = MagicMock()
        cell_mock.value = "Original"
        ws.__getitem__.return_value = cell_mock

        exporter.fill_holding_information(ws)

        mp = default_stallkarte_holding_mapping
        ws.__setitem__.assert_any_call(mp.animals_count, 1234)

    def test_fill_checklist(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)

        exporter.fill_checklist(ws)

        mp = default_stallkarte_checklist_mapping

        ws.__setitem__.assert_any_call(mp.stable_disinfection_date, "10.01.2023")
        ws.__setitem__.assert_any_call(mp.stable_disinfectant, "Disinfectant A")
        ws.__setitem__.assert_any_call(mp.silo_cleaning_date, "11.01.2023")
        ws.__setitem__.assert_any_call(mp.water_line_disinfection_date, "12.01.2023")
        ws.__setitem__.assert_any_call(
            mp.pest_control_measures_annotation, "Pest control note"
        )

    def test_fill_production_day(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)

        day = MagicMock(spec=StallkarteStateDay)
        day.production_day = 0
        day.temperature_celsius = 25.0
        day.humidity_percent = 60.0
        day.water_consumption_liters = 100
        day.feed_consumption_kg = 50
        day.weight_grams = 200
        day.notes = [
            NoteEntry(
                id="note-1",
                note_type=NoteType.OTHER,
                note_text="Day note",
            )
        ]

        section_data = MagicMock()
        section_data.natural_mortality_morning = 1
        section_data.natural_mortality_evening = 0
        section_data.selective_mortality_morning = 0
        section_data.selective_mortality_evening = 1
        section_data.did_inspection_morning = True
        section_data.did_inspection_evening = False
        section_data.note = "Section note"

        day.sections = {1: section_data}

        exporter.fill_production_day(ws, day)

        mp = get_mapping_of_production_day(0)

        ws.__setitem__.assert_any_call(mp.day, 0)
        ws.__setitem__.assert_any_call(mp.temperature, 25.0)
        ws.__setitem__.assert_any_call(mp.humidity, 0.6)
        ws.__setitem__.assert_any_call(mp.water_consumption, 100)
        ws.__setitem__.assert_any_call(mp.feed_consumption, 50)
        ws.__setitem__.assert_any_call(mp.weight, 200)
        ws.__setitem__.assert_any_call(mp.natural_mortality, 1)
        ws.__setitem__.assert_any_call(mp.selective_mortality, 1)
        ws.__setitem__.assert_any_call(mp.did_inspection_1, "X")
        ws.__setitem__.assert_any_call(mp.did_inspection_2, "")
        ws.__setitem__.assert_any_call(
            mp.annotation,
            "Sonstiges - Day note",
        )

    def test_fill_production_day_formats_typed_notes(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)

        day = MagicMock(spec=StallkarteStateDay)
        day.production_day = 0
        day.temperature_celsius = None
        day.humidity_percent = None
        day.water_consumption_liters = None
        day.feed_consumption_kg = None
        day.weight_grams = None
        day.notes = [
            NoteEntry(
                id="vacc-1",
                note_type=NoteType.VACCINATION,
                vaccination_code=VaccinationCode.ND,
                note_text="Planimpfung",
            ),
            NoteEntry(
                id="treat-1",
                note_type=NoteType.TREATMENT,
                treatment_code=TreatmentCode.AMPROLINE,
                treatment_amount_value=1.5,
                treatment_amount_unit=TreatmentAmountUnit.L_PER_1000,
                treatment_waiting_time_value=2,
                treatment_waiting_time_unit=WaitingTimeUnit.DAY,
                note_text="Dokumentiert",
            ),
            NoteEntry(
                id="sock-1",
                note_type=NoteType.SOCK_TEST,
                sock_test_result=SockTestResult.NEGATIVE,
            ),
        ]
        day.sections = {1: MagicMock(note=None)}

        exporter.fill_production_day(ws, day)

        mp = get_mapping_of_production_day(0)
        ws.__setitem__.assert_any_call(
            mp.annotation,
            "Impfung - ND - Planimpfung / "
            "Behandlung - amproline - 1.5 l/1000 - Wartezeit 2 day - Dokumentiert / "
            "Sockenprobe - negative",
        )

    def test_fill_production_days(
        self, mock_holding: MagicMock, mock_stallkarte: MagicMock
    ) -> None:
        exporter = StallkarteWorksheetExporter(mock_holding, mock_stallkarte, 1)
        ws = MagicMock(spec=Worksheet)

        day1 = MagicMock(spec=StallkarteStateDay)
        day2 = MagicMock(spec=StallkarteStateDay)
        mock_stallkarte.state.days = {1: day1, 2: day2}

        with patch.object(exporter, "fill_production_day") as mock_fill_day:
            exporter.fill_production_days(ws)

            assert mock_fill_day.call_count == 2
            mock_fill_day.assert_any_call(ws, day1)
            mock_fill_day.assert_any_call(ws, day2)
